"""Build presentation/decision_engine.pptx — 10 slides."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
BG = (18, 32, 48)
ACCENT = (0, 180, 160)
WHITE = (255, 255, 255)
GRAY = (200, 205, 212)


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(*BG)


def add(slide, left, top, width, height, text, size=20, bold=False,
        color=WHITE, align=PP_ALIGN.LEFT):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top),
                                  Inches(width), Inches(height)).text_frame
    tx.word_wrap = True
    p = tx.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(*color)
    p.alignment = align
    return tx


def bullets(slide, left, top, width, height, items, size=18):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top),
                                  Inches(width), Inches(height)).text_frame
    tx.word_wrap = True
    for i, it in enumerate(items):
        p = tx.add_paragraph() if i else tx.paragraphs[0]
        p.text = it
        p.font.size = Pt(size)
        p.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(*GRAY)
        p.level = 0
        p.space_after = Pt(8)


SLIDES = [
    ("Predictive Maintenance:\nthe Decision Engine",
     ["Project #3 — the layer between prediction and action",
      "Mukesh Sharma · github.com/designmukesh",
      "Builds on: hydraulic fault classification (#1) + turbofan RUL (#2)"]),
    ("The gap this closes",
     ["Models answer 'what is the condition?' — planners need 'what do we do Tuesday?'",
      "Most pilots die at the handoff: accurate model, no decision policy",
      "This project engineers the handoff: costs → rule → backtested policy"]),
    ("Cost model (assumed, labeled)",
     ["planned intervention = 1.0 (numeraire)",
      "unplanned failure = k × planned, k ∈ {3, 5, 10} scenarios",
      "inspection = 0.1 · lead time = 1 cycle",
      "Only the ratio k matters for the decision"]),
    ("The decision rule",
     ["Act when E[cost of waiting] > E[cost of acting]",
      "p(fail) × C_u > C_p  ⟺  p > 1/k",
      "p from 3,524 OBSERVED residuals — no assumed Gaussian",
      "Thresholds from economics, not folklore"]),
    ("Backtest design",
     ["20 validation engines, out-of-sample RUL predictions",
      "Full run-to-failure trajectories — true failure known, bills exact",
      "Three policies: run-to-failure · fixed-interval (N) · predictive (T)",
      "Sweep T 0–60; sensitivity over k = 3 / 5 / 10"]),
    ("Results (k=5, illustrative)",
     ["Run-to-failure: 5.0/engine, 20/20 unplanned failures",
      "Fixed-interval N=25: 1.0/engine",
      "Predictive T*=8: 1.0/engine, 0 unplanned, least wasted life",
      "Ranking holds across all k scenarios"]),
    ("Two honest surprises",
     ["Cost curve is L-shaped, not U-shaped — wasted life isn't monetized",
      "T* doesn't move with k here: k prices timidity, not the optimum",
      "Selection is lexicographic: min cost, then min wasted life",
      "Flat bottom = threshold error is cheap (good for deployment)"]),
    ("Hydraulic fault policy",
     ["Auditable table: (component, fault state) → operate/monitor/inspect/replace",
      "1,764 held-out predictions: expected action cost 0.295/cycle",
      "1 missed severe state in 380 — the residual risk, stated plainly"]),
    ("Deploying it",
     ["Policy service: predict → decide → work order",
      "Re-estimate the residual distribution as field data arrives",
      "Monitor two vital signs: slip rate + wasted healthy life",
      "Next step: field pilot with holdout fleet — not a rollout"]),
    ("Takeaway",
     ["The method is the deliverable — any planner's costs plug in",
      "14 tests · 2 executed notebooks · interactive Streamlit app",
      "Limitations volunteered: assumed costs, n=20, single failure mode",
      "Questions?"]),
]

for title, pts in SLIDES:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    add(slide, 0.7, 0.4, 11.9, 1.6, title, size=40, bold=True)
    bullets(slide, 0.9, 2.2, 11.5, 4.8, ["•  " + p for p in pts], size=22)

prs.save("../presentation/decision_engine.pptx")
print("wrote ../presentation/decision_engine.pptx")
